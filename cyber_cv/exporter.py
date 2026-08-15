from __future__ import annotations

import html
import io
import json
import pathlib
import re
import zipfile

from . import TITLE, WATERMARK


def normalize(raw: dict) -> dict:
    profile = raw.get("profile") or {}
    entries = []
    for index, item in enumerate(raw.get("entries") or []):
        text = str(item.get("rawInput") or "").strip()
        if not text:
            continue
        entries.append({
            "order": int(item.get("order", index)),
            "timestamp": str(item.get("timestamp") or ""),
            "workspace": str(item.get("workspace") or "Personal"),
            "organization": str(item.get("organization") or ""),
            "source": str(item.get("source") or ""),
            "rawInput": text,
            "annotation": str(item.get("annotation") or "").strip(),
            "recordId": str(item.get("recordId") or ""),
            "verified": bool(item.get("verified", False)),
        })
    entries.sort(key=lambda item: (item["order"], item["timestamp"]))
    return {
        "title": TITLE,
        "displayName": str(profile.get("displayName") or "Prompt Trace User"),
        "authorId": str(profile.get("authorId") or "").upper(),
        "headline": str(profile.get("headline") or ""),
        "summary": str(profile.get("summary") or ""),
        "entries": entries,
    }


def render_markdown(data: dict) -> str:
    lines = [f"# {TITLE}", "", f"## {data['displayName']} · [PT:{data['authorId']}]", ""]
    if data["headline"]:
        lines += [data["headline"], ""]
    if data["summary"]:
        lines += [data["summary"], ""]
    lines += ["## Selected prompt and terminal inputs", ""]
    for item in data["entries"]:
        context = item["organization"] or item["workspace"]
        verified = " · Verified" if item["verified"] else ""
        lines += [f"### {item['timestamp']} · {context}{verified}", "", f"> {item['rawInput'].replace(chr(10), chr(10) + '> ')}", ""]
        if item["annotation"]:
            lines += [f"**Context:** {item['annotation']}", ""]
        lines += [f"Source: `{item['source']}` · Record: `{item['recordId']}`", ""]
    lines += ["---", f"*{WATERMARK}*", ""]
    return "\n".join(lines)


def render_html(data: dict) -> str:
    esc = html.escape
    entries = []
    for item in data["entries"]:
        context = item["organization"] or item["workspace"]
        entries.append(f'''<article class="entry"><div class="meta">{esc(item["timestamp"])} · {esc(context)}{' · Verified' if item['verified'] else ''}</div><pre>{esc(item["rawInput"])}</pre>{f'<p class="annotation">{esc(item["annotation"])}</p>' if item['annotation'] else ''}<div class="record">{esc(item["source"])} · {esc(item["recordId"])}</div></article>''')
    return f'''<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{TITLE}</title><style>@page{{margin:18mm}}:root{{--pink:#ff2da1;--ink:#211720;--muted:#715f70}}*{{box-sizing:border-box}}body{{max-width:900px;margin:40px auto;padding:0 28px;color:var(--ink);font:15px/1.55 system-ui,sans-serif}}header{{border-bottom:3px solid var(--pink);padding-bottom:22px}}h1{{margin:0;font-size:30px}}.identity{{color:var(--pink);font:700 16px ui-monospace,monospace}}.entry{{break-inside:avoid;border-bottom:1px solid #eadde7;padding:22px 0}}.meta,.record{{color:var(--muted);font-size:12px}}pre{{white-space:pre-wrap;font:15px/1.55 ui-monospace,monospace;background:#faf6f9;border-left:4px solid var(--pink);padding:14px}}.annotation{{font-style:italic}}.watermark{{position:fixed;right:16px;bottom:12px;color:rgba(255,45,161,.55);font:700 11px system-ui,sans-serif;letter-spacing:.08em}}</style></head><body><header><h1>{TITLE}</h1><div class="identity">{esc(data['displayName'])} · [PT:{esc(data['authorId'])}]</div><h2>{esc(data['headline'])}</h2><p>{esc(data['summary'])}</p></header><main><h2>Selected prompt and terminal inputs</h2>{''.join(entries)}</main><div class="watermark">{WATERMARK}</div></body></html>'''


def render_pdf(data: dict) -> bytes:
    try:
        from weasyprint import HTML
        return HTML(string=render_html(data)).write_pdf()
    except ImportError:
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import LETTER
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        except ImportError as error:
            raise RuntimeError("Portable Document Format export requires WeasyPrint or ReportLab") from error
        stream = io.BytesIO()
        styles = getSampleStyleSheet()
        prompt_style = ParagraphStyle("Prompt", parent=styles["BodyText"], leftIndent=12, borderColor=colors.HexColor("#ff2da1"), borderWidth=1, borderPadding=8, spaceAfter=8)
        document = SimpleDocTemplate(stream, pagesize=LETTER, title=TITLE)
        story = [Paragraph(TITLE, styles["Title"]), Paragraph(f"{html.escape(data['displayName'])} · [PT:{html.escape(data['authorId'])}]", styles["Heading2"])]
        if data["headline"]: story.append(Paragraph(html.escape(data["headline"]), styles["Heading3"]))
        if data["summary"]: story.append(Paragraph(html.escape(data["summary"]), styles["BodyText"]))
        story.append(Spacer(1, 14))
        for item in data["entries"]:
            context = item["organization"] or item["workspace"]
            story.append(Paragraph(f"{html.escape(item['timestamp'])} · {html.escape(context)}", styles["Heading3"]))
            story.append(Paragraph(html.escape(item["rawInput"]).replace("\n", "<br/>"), prompt_style))
            if item["annotation"]: story.append(Paragraph(f"Context: {html.escape(item['annotation'])}", styles["BodyText"]))
        watermark_style = ParagraphStyle("Watermark", parent=styles["BodyText"], fontSize=8, textColor=colors.HexColor("#ff2da1"))
        story.append(Spacer(1, 16)); story.append(Paragraph(WATERMARK, watermark_style))
        document.build(story); return stream.getvalue()


def render_docx(data: dict) -> bytes:
    try:
        from docx import Document
        from docx.shared import RGBColor, Pt
    except ImportError as error:
        raise RuntimeError("Microsoft Word export requires python-docx") from error
    document = Document()
    document.add_heading(TITLE, 0)
    identity = document.add_paragraph(f"{data['displayName']} · [PT:{data['authorId']}]")
    identity.runs[0].font.color.rgb = RGBColor(255, 45, 161)
    if data["headline"]:
        document.add_heading(data["headline"], 1)
    if data["summary"]:
        document.add_paragraph(data["summary"])
    document.add_heading("Selected prompt and terminal inputs", 1)
    for item in data["entries"]:
        context = item["organization"] or item["workspace"]
        document.add_heading(f"{item['timestamp']} · {context}", 2)
        prompt = document.add_paragraph(item["rawInput"])
        prompt.style = document.styles["Quote"]
        if item["annotation"]:
            document.add_paragraph(f"Context: {item['annotation']}")
        document.add_paragraph(f"{item['source']} · {item['recordId']}")
    footer = document.sections[0].footer.paragraphs[0]
    footer.text = WATERMARK
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = RGBColor(255, 45, 161)
    stream = io.BytesIO(); document.save(stream); return stream.getvalue()


def export_bundle(raw: dict, destination: pathlib.Path) -> pathlib.Path:
    data = normalize(raw)
    slug = re.sub(r"[^A-Za-z0-9]+", "-", data["authorId"] or "user").strip("-").lower()
    markdown = render_markdown(data)
    html_document = render_html(data)
    destination.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(destination, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(f"{slug}-terminal-prompt-trace-cv.md", markdown)
        archive.writestr(f"{slug}-terminal-prompt-trace-cv.html", html_document)
        archive.writestr(f"{slug}-terminal-prompt-trace-cv.pdf", render_pdf(data))
        archive.writestr(f"{slug}-terminal-prompt-trace-cv.docx", render_docx(data))
        archive.writestr("manifest.json", json.dumps({"title": TITLE, "generator": WATERMARK, "formats": ["markdown", "html", "pdf", "docx"]}, indent=2))
    return destination
